# -*- coding: utf-8 -*-
"""
consultas.py — guarda e lê a planilha organizada e responde as perguntas da Maria
com regras simples (funciona sem internet e sem IA).
"""
import re
from datetime import date, timedelta
import pandas as pd
from openpyxl import load_workbook

from maria import COLUNAS, norm, MESES

ARQUIVO_PADRAO = "Vendas da Maria - organizada.xlsx"


def eh_data(d):
    return isinstance(d, date) and pd.notna(d)


def dinheiro(v):
    try:
        return ("R$ " + f"{float(v):,.2f}").replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return "R$ ?"


def carregar(caminho=ARQUIVO_PADRAO):
    df = pd.read_excel(caminho, sheet_name="Vendas")
    df["Data"] = pd.to_datetime(df["Data"], errors="coerce").dt.date
    for c in ("Valor (R$)", "Frete (R$)", "Parcelas"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    for c in ("Produto", "Cliente", "Pagamento", "Cartão", "Entrega", "Rastreio", "Recebido?", "Obs", "Confirmar?", "Nota do sistema"):
        df[c] = df[c].fillna("").astype(str)
    return df


def gravar(venda, caminho=ARQUIVO_PADRAO):
    """Acrescenta a venda na planilha (aba Vendas, ou 'Não é venda'). Devolve o nº da linha."""
    wb = load_workbook(caminho)
    if venda.get("Tipo", "Venda") == "Venda":
        ws = wb["Vendas"]
        cols = COLUNAS
    else:
        ws = wb["Não é venda"]
        cols = ["Data", "Tipo", "Produto", "Cliente", "Valor (R$)", "Pagamento", "Obs", "Nota do sistema", "Nº original"]
    linha = ws.max_row + 1
    while linha > 2 and all(ws.cell(row=linha - 1, column=j).value in (None, "") for j in range(1, len(cols) + 1)):
        linha -= 1  # pula linhas vazias deixadas pelo Excel
    for j, nome in enumerate(cols, 1):
        v = venda.get(nome)
        if nome == "Nº original":
            v = "app"
        if isinstance(v, float) and v != v:
            v = None
        c = ws.cell(row=linha, column=j, value=v)
        if nome == "Data":
            c.number_format = "dd/mm/yyyy"
        if nome in ("Valor (R$)", "Frete (R$)"):
            c.number_format = '"R$" #,##0.00'
    if ws.title == "Vendas":
        ws.auto_filter.ref = f"A1:O{linha}"
    wb.save(caminho)
    return linha


def marcar_recebido(cliente, caminho=ARQUIVO_PADRAO):
    """'a Rita me pagou' -> Recebido? = Sim nas vendas fiadas dela."""
    wb = load_workbook(caminho)
    ws = wb["Vendas"]
    n = 0
    for r in range(2, ws.max_row + 1):
        if norm(ws.cell(row=r, column=3).value) == norm(cliente) and norm(ws.cell(row=r, column=11).value) == "nao":
            ws.cell(row=r, column=11, value="Sim")
            n += 1
    wb.save(caminho)
    return n


# ---------------------------------------------------------------------------
# Perguntas
# ---------------------------------------------------------------------------
INTERROGATIVAS = r"\b(quanto|quantas|quantos|quem|qual|quais|cade|onde|quando|como foi|lista|mostra|me fala|me diz|quero saber|tem algum|tem alguma|ja pagou|chegou|resumo|total|vendi|faturei|devem|deve|devendo|fiado|correio|rastreio|encomenda|ultima venda|ultimas vendas)\b"


def parece_pergunta(texto):
    t = norm(texto)
    if "?" in texto:
        return True
    if re.search(r"\b(vendi|vendeu|saiu|comprei|devolveu|mandei|enviei)\b\s+(um|uma|o|a|\d|r\$)", t):
        return False  # "vendi um colar..." é venda
    if re.match(r"^(quanto|quem|qual|quais|cade|onde|quando|lista|mostra|me fala|me diz|quero saber|tem)\b", t):
        return True
    if not re.search(r"\d", t) and re.search(INTERROGATIVAS, t):
        return True
    return False


def _periodo(t, hoje):
    """Devolve (inicio, fim, rotulo) ou None."""
    if re.search(r"\b(hoje|hj)\b", t):
        return hoje, hoje, "hoje"
    if re.search(r"\bontem\b", t):
        d = hoje - timedelta(days=1)
        return d, d, "ontem"
    if re.search(r"\b(semana|7 dias)\b", t):
        return hoje - timedelta(days=6), hoje, "nos últimos 7 dias"
    for nome, m in MESES.items():
        if len(nome) > 3 and re.search(r"\b" + nome + r"\b", t):
            ano = hoje.year if m <= hoje.month else hoje.year - 1
            ini = date(ano, m, 1)
            fim = (date(ano + (m == 12), (m % 12) + 1, 1) - timedelta(days=1))
            return ini, fim, f"em {nome}"
    if re.search(r"\b(mes|neste mes|esse mes|no mes)\b", t):
        ini = date(hoje.year, hoje.month, 1)
        return ini, hoje, "neste mês"
    if re.search(r"\b(mes passado)\b", t):
        fim = date(hoje.year, hoje.month, 1) - timedelta(days=1)
        return date(fim.year, fim.month, 1), fim, "no mês passado"
    if re.search(r"\b(total|tudo|ate agora|no geral|desde o comeco)\b", t):
        return None, None, "no total"
    return None


def _acha_cliente(t, df):
    nomes = sorted({c for c in df["Cliente"].unique() if c and c not in ("Não informado", "Não identificado")},
                   key=len, reverse=True)
    for nome in nomes:
        n = norm(nome)
        if re.search(r"\b" + re.escape(n) + r"\b", t):
            return nome
        # "da Ana", "pra Ana", "a Ana" -> só o primeiro nome
        primeiro = n.split()[0] if n.split() else n
        if len(primeiro) > 2 and primeiro not in ("tia", "seu", "menina", "a", "o") and re.search(r"\b" + re.escape(primeiro) + r"\b", t):
            return nome
    return None


def _acha_produto(t, df):
    prods = sorted({p for p in df["Produto"].unique() if p}, key=len, reverse=True)
    for p in prods:
        if re.search(r"\b" + re.escape(norm(p)) + r"s?\b", t):
            return p
    return None


def _linhas(dfx, maximo=8):
    out = []
    for _, r in dfx.sort_values("Data", ascending=False).head(maximo).iterrows():
        d = r["Data"].strftime("%d/%m") if eh_data(r["Data"]) else "sem data"
        pag = r["Pagamento"] + (f" {int(r['Parcelas'])}x" if r["Parcelas"] and r["Parcelas"] > 1 else "")
        extra = f" · rastreio {r['Rastreio']}" if r["Rastreio"] else ""
        pend = " · **ainda não pagou**" if r["Recebido?"] == "Não" else ""
        out.append(f"- {d} — {r['Produto']} pra {r['Cliente']}, {dinheiro(r['Valor (R$)'])}, {pag}{extra}{pend}")
    return "\n".join(out)


def responder(pergunta, df, hoje=None):
    hoje = hoje or date.today()
    t = norm(pergunta)
    vendas = df[df["Produto"] != ""]

    # quem me deve
    if re.search(r"\b(deve|devem|devendo|fiado|receber|a receber|me pagar|nao pagou|pendente)\b", t) and not re.search(r"\bme pagou\b", t):
        dev = vendas[vendas["Recebido?"] == "Não"]
        if dev.empty:
            return "Ninguém está te devendo. Tudo recebido. 🎉"
        total = dev["Valor (R$)"].sum()
        linhas = []
        for nome, g in dev.groupby("Cliente"):
            datas = [x for x in g["Data"] if eh_data(x)]
            d = min(datas) if datas else None
            desde = f" (desde {d.strftime('%d/%m')})" if d else ""
            itens = ", ".join(g["Produto"])
            linhas.append(f"- **{nome}**: {dinheiro(g['Valor (R$)'].sum())}{desde} — {itens}")
        rodape = ""
        if (dev["Cliente"] == "Não identificado").any():
            rodape = "\n\nTem um fiado sem nome: você anotou como \"não lembro\". Vale tentar lembrar quem foi."
        return f"Estão te devendo **{dinheiro(total)}** no total:\n\n" + "\n".join(linhas) + rodape + \
            "\n\nQuando alguém pagar, é só me dizer: \"a Rita me pagou\"."

    # "a Rita me pagou"
    m = re.search(r"\b(?:a|o)?\s*([a-z]+(?:\s[a-z]+)?)\s+me pagou\b", t)
    if m:
        cli = _acha_cliente(t, df)
        if cli:
            return {"acao": "receber", "cliente": cli}
        return "Não achei essa pessoa nas vendas fiadas. Quem te pagou?"

    # correio / rastreio
    if re.search(r"\b(correio|correios|rastreio|rastreamento|codigo|encomenda|encomendas|sedex|enviei|mandei|postei|envio|envios|motoboy)\b", t):
        env = vendas[vendas["Entrega"].isin(["Correios", "Motoboy"])]
        cli = _acha_cliente(t, df)
        if cli:
            e = env[env["Cliente"] == cli]
            if e.empty:
                return f"Não achei envio pra **{cli}**. As compras dela foram retiradas na loja ou estão sem entrega anotada."
            cod = e[e["Rastreio"] != ""]
            txt = f"Envios pra **{cli}**:\n\n" + _linhas(e)
            if cod.empty:
                txt += "\n\nNenhum deles tem código de rastreio anotado."
            return txt
        com = env[env["Rastreio"] != ""]
        sem = env[(env["Rastreio"] == "") & (env["Entrega"] == "Correios")]
        txt = "**Encomendas com código de rastreio:**\n\n" + (_linhas(com, 12) if not com.empty else "- nenhuma")
        if not sem.empty:
            txt += f"\n\nMais {len(sem)} envios pelos Correios estão **sem código anotado**. Da próxima vez me manda o código que eu guardo."
        return txt

    # última venda
    if re.search(r"\b(ultima venda|ultimas vendas|o que vendi por ultimo)\b", t):
        return "Últimas vendas:\n\n" + _linhas(vendas, 5)

    # produto mais vendido
    if re.search(r"\b(mais vend\w*|mais sai\w*|campe\w*|melhor produto|top produto)", t):
        top = vendas.groupby("Produto")["Valor (R$)"].agg(["count", "sum"]).sort_values("count", ascending=False).head(5)
        linhas = [f"- **{p}**: {int(r['count'])} vendas, {dinheiro(r['sum'])}" for p, r in top.iterrows()]
        return "Os que mais saem:\n\n" + "\n".join(linhas)

    # cliente específico
    cli = _acha_cliente(t, df)
    prod = _acha_produto(t, df)
    per = _periodo(t, hoje)

    if cli and not re.search(r"\b(quanto vendi|faturei)\b", t):
        c = vendas[vendas["Cliente"] == cli]
        total = c["Valor (R$)"].sum()
        pend = c[c["Recebido?"] == "Não"]["Valor (R$)"].sum()
        txt = f"**{cli}** comprou {len(c)} vez(es), {dinheiro(total)} no total:\n\n" + _linhas(c)
        if pend:
            txt += f"\n\n⚠️ Está devendo {dinheiro(pend)}."
        return txt

    if prod and re.search(r"\b(quem comprou|quem levou|quem pegou)\b", t):
        p = vendas[vendas["Produto"] == prod]
        return f"Quem comprou **{prod}** ({len(p)} vendas):\n\n" + _linhas(p, 12)

    # quanto vendi (período)
    if per or re.search(r"\b(quanto vendi|quanto faturei|quanto fiz|quanto deu|vendas de|vendi)\b", t):
        if per is None:
            per = (date(hoje.year, hoje.month, 1), hoje, "neste mês")
        ini, fim, rotulo = per
        if ini is None:
            sel = vendas
        else:
            sel = vendas[[eh_data(d) and ini <= d <= fim for d in vendas["Data"]]]
        if prod:
            sel = sel[sel["Produto"] == prod]
        if sel.empty:
            return f"Nenhuma venda {rotulo}" + (f" de {prod}" if prod else "") + "."
        total = sel["Valor (R$)"].sum()
        por_pag = sel.groupby("Pagamento")["Valor (R$)"].sum().sort_values(ascending=False)
        detalhe = ", ".join(f"{k} {dinheiro(v)}" for k, v in por_pag.items())
        txt = f"Você vendeu **{dinheiro(total)}** {rotulo} em {len(sel)} venda(s)" + (f" de {prod}" if prod else "") + f".\n\nPor forma de pagamento: {detalhe}."
        if len(sel) <= 8:
            txt += "\n\n" + _linhas(sel)
        return txt

    return ("Não entendi direito. Você pode me perguntar, por exemplo:\n\n"
            "- *quem me deve?*\n- *o que está no correio?*\n- *quanto vendi hoje?* / *no mês?* / *em setembro?*\n"
            "- *o que a Ana comprou?*\n\nOu me contar uma venda: *vendi um colar pra Ana, 89 no pix, mandei pelo correio*.")
